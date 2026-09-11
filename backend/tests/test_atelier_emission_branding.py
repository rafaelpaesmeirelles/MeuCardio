"""Brand-only format checks; no patient data, signing or external delivery."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from docx import Document
from docx.shared import Inches as DocxInches
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader

from app.services.pdf.marca import LOGO
from app.services.pdf.nucleo import ler_png

ROOT = Path(__file__).resolve().parents[2]


def test_atelier_raster_is_shared_print_safe_and_decodable_by_both_pdf_engines():
    public_logo = ROOT / "frontend/public/atelier/corvia-logo-atelier.png"
    assert LOGO.name == "corvia-logo-atelier.png"
    assert LOGO.read_bytes() == public_logo.read_bytes()
    with Image.open(LOGO) as logo:
        assert logo.size == (1800, 480)
        assert logo.convert("RGB").getpixel((0, 0)) == (255, 255, 255)
    assert ImageReader(str(LOGO)).getSize() == (1800, 480)
    decoded = ler_png(str(LOGO))
    assert (decoded["largura"], decoded["altura"]) == (1800, 480)


def test_atelier_logo_embeds_unchanged_in_word_and_powerpoint():
    word = Document()
    word.add_picture(str(LOGO), width=DocxInches(1.5))
    word_bytes = BytesIO()
    word.save(word_bytes)

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_picture(str(LOGO), Inches(0.5), Inches(0.5), width=Inches(2.5))
    slide_bytes = BytesIO()
    slides.save(slide_bytes)

    expected = LOGO.read_bytes()
    for stream, media_prefix in ((word_bytes, "word/media/"), (slide_bytes, "ppt/media/")):
        with ZipFile(stream) as archive:
            assert any(
                archive.read(name) == expected
                for name in archive.namelist()
                if name.startswith(media_prefix)
            )


def test_atelier_maskable_marks_reserve_twenty_percent_padding():
    for size in (192, 512):
        with Image.open(ROOT / f"frontend/public/atelier/corvia-mark-atelier-{size}.png") as mark:
            assert mark.size == (size, size)
            rgb = mark.convert("RGB")
            padding = (size + 4) // 5
            for point in ((0, 0), (size // 2, padding - 1), (padding - 1, size // 2)):
                assert rgb.getpixel(point) == (255, 253, 245)
