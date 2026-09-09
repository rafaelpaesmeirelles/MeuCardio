"""Exercise PDF, Word and slides without a database or real patient records."""
from io import BytesIO
from types import SimpleNamespace

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
from pptx import Presentation

from app.services import exportacao_conteudo, exportacao_office, material_paciente


def _professional():
    return SimpleNamespace(
        full_name="Profissional de demonstração", professional_title="Dra.", profession="Médica",
        council_name="CRM", council_state="SP", council_number="000000", rqe="000000",
        specialty="Cardiologia", document_logo_url=None, include_workplace_on_documents=True,
        workplace_name="Clínica de demonstração", workplace_department="Cardiologia",
        workplace_role="Médica", workplace_notes="",
    )


def _content():
    return [exportacao_conteudo.ConteudoExportavel(
        tipo="material_paciente", slug="teste", titulo="Material de demonstração", tema="Teste de diagramação",
        secoes=[exportacao_conteudo.SecaoExportacao(
            titulo="Orientações", paragrafos=["PARAGRAFOTESTE " * 240],
            itens=["ITEMTESTE " * 80], destaque="DESTAQUETESTE " * 180,
        )],
    )]


def test_pdf_export_preserves_long_sections_and_repeats_identity():
    data = exportacao_conteudo.gerar_pdf(_content(), user=_professional(), incluir_dados_assinante=True)
    with pdfplumber.open(BytesIO(data)) as pdf:
        text = " ".join(p.extract_text() for p in pdf.pages)
        for token, count in [("PARAGRAFOTESTE", 240), ("ITEMTESTE", 80), ("DESTAQUETESTE", 180)]:
            assert text.count(token) == count
        for page in pdf.pages:
            assert "Profissional" in page.extract_text()
            assert all(61 <= c["x0"] <= c["x1"] <= page.width - 61 for c in page.chars)
            body = [w for w in page.extract_words() if w["text"].endswith("TESTE")]
            assert all(w["bottom"] < page.height - 76 for w in body)


def test_signed_pdf_export_reserves_the_stamp_area_on_every_page():
    data = exportacao_conteudo.gerar_pdf(_content(), user=_professional(),
                                        incluir_dados_assinante=True, reservar_assinatura=True)
    with pdfplumber.open(BytesIO(data)) as pdf:
        assert len(pdf.pages) > 1
        for page in pdf.pages:
            assert all(c["bottom"] < page.height - 88 for c in page.chars)


def test_word_export_has_centered_logo_and_right_identity_in_repeating_header(monkeypatch, tmp_path):
    logo = tmp_path / "logo.png"
    Image.new("RGB", (280, 120), (32, 74, 100)).save(logo)
    monkeypatch.setattr(exportacao_office, "rendered_logo_png", lambda _url: logo)
    data = exportacao_office.gerar_docx(_content(), user=_professional(), incluir_dados_assinante=True)
    doc = Document(BytesIO(data))
    table = doc.sections[0].header.tables[0]
    assert table.cell(0, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.cell(0, 2).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert "Profissional" in table.cell(0, 2).text
    assert table.columns[0].width == table.columns[2].width
    assert len(table.cell(0, 1)._tc.xpath('.//w:drawing')) == 1
    text = " ".join(p.text for p in doc.paragraphs)
    for token, count in [("PARAGRAFOTESTE", 240), ("ITEMTESTE", 80), ("DESTAQUETESTE", 180)]:
        assert text.count(token) == count


def test_pptx_export_preserves_all_patient_guidance_as_editable_text():
    data = exportacao_office.gerar_pptx(_content(), user=_professional(), incluir_dados_assinante=True)
    pptx = Presentation(BytesIO(data))
    text = " ".join(shape.text for slide in pptx.slides for shape in slide.shapes if shape.has_text_frame)
    for token, count in [("PARAGRAFOTESTE", 240), ("ITEMTESTE", 80), ("DESTAQUETESTE", 180)]:
        assert text.count(token) == count
